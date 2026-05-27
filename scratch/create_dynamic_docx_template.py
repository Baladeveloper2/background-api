import docx

src_path = r"d:\project\frontend\public\assets\CL-2026-05-0010 - Integra Software (1).docx"
dest_path = r"d:\project\frontend\public\assets\invoice_template.docx"

doc = docx.Document(src_path)
table = doc.tables[0]

print("Modifying cells to map to dynamic placeholders...")

# 1. Update Period Header in Row 0
# Cell 0 has "Bill copy for the month of May-2026"
row0 = table.rows[0]
for cell in row0.cells:
    cell.text = "Bill copy for the period {{billingPeriod}}"

# 2. Update Invoice Meta in Row 3
row3 = table.rows[3]
# Set Invoice No
row3.cells[0].text = "{{invoiceNo}}"
row3.cells[1].text = "{{invoiceNo}}"
# Set Invoice Date
row3.cells[2].text = "{{invoiceDate}}"
row3.cells[3].text = "{{invoiceDate}}"
# Set Due Date
row3.cells[4].text = "{{dueDate}}"
row3.cells[5].text = "{{dueDate}}"

# 3. Update Client Info in Row 5
row5 = table.rows[5]
client_info_template = "{{clientName}}\nAddress: {{clientAddress}}\nGST: {{clientGst}}"
for i in range(1, 6):
    row5.cells[i].text = client_info_template

# 4. Update Details Table Header in Row 7
row7 = table.rows[7]
row7.cells[0].text = "S.No"
row7.cells[1].text = "Candidate Name"
row7.cells[2].text = "Case Ref"
row7.cells[3].text = "Module / Check Type"
row7.cells[4].text = "Completion Date"
row7.cells[5].text = "Amount"

# 5. Update Row 8 to be the dynamic loop row
row8 = table.rows[8]
row8.cells[0].text = "{#billingRows}{{sNo}}"
row8.cells[1].text = "{{candidateName}}"
row8.cells[2].text = "{{caseRef}}"
row8.cells[3].text = "{{moduleName}}"
row8.cells[4].text = "{{completionDate}}"
row8.cells[5].text = "₹{{amount}}{/billingRows}"

# 6. Delete Row 9 (which is the extra static row)
# To delete a row in python-docx:
# We access the table element and remove the XML element representing row 9
tbl_el = table._tbl
row9_el = table.rows[9]._tr
tbl_el.remove(row9_el)
print("Deleted duplicate static row 9.")

# 7. Update Subtotal in Row 9 (originally Row 10)
# Let's re-verify row index after deleting row 9:
# Original Row 10 is now Row 9
row9 = table.rows[9]
row9.cells[5].text = "₹{{subtotal}}"

# 8. Update IGST in Row 10 (originally Row 11)
row10 = table.rows[10]
row10.cells[3].text = "{{gstPercentage}}%"
row10.cells[4].text = "{{gstPercentage}}%"
row10.cells[5].text = "₹{{gst}}"

# 9. Update Grand Total in Row 12 (originally Row 13)
row12 = table.rows[12]
row12.cells[5].text = "₹{{grandTotal}}"

# 10. Update Total Amount in Words in Row 13 (originally Row 14)
row13 = table.rows[13]
for cell in row13.cells:
    cell.text = "Total Amount (In Words): {{grandTotalInWords}}\nPayment Information"

doc.save(dest_path)
print("Successfully generated dynamic docx template at:", dest_path)
